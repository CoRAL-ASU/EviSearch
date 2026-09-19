"""Read a table spreadsheet: one header row (an optional section row above it is skipped) and any filled rows.

A row may name its paper in a "Document Name" column. Blank rows, system columns (`_SYS`) and unnamed columns are
dropped; page notes such as "(pg3, text)" are stripped from cell values.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

DOC_COLUMN = "Document Name"
PAGE_NOTE_RE = re.compile(r"\s*\((?:pg|page)\s*\d+[^)]*\)", re.I)


@dataclass
class Sheet:
    headers: List[str]
    rows: List[Dict[str, str]] = field(default_factory=list)  # header -> cleaned value ("" when empty)

    def row_for(self, doc_hint: str) -> Optional[Dict[str, str]]:
        """The row whose Document Name contains doc_hint (case-insensitive, ignoring the .pdf extension)."""
        hint = doc_hint.lower().removesuffix(".pdf")
        for row in self.rows:
            if hint and hint in row.get(DOC_COLUMN, "").lower():
                return row
        return None


def clean_value(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = PAGE_NOTE_RE.sub("", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _is_section_row(cells: List[Any]) -> bool:
    """A section row labels groups of columns: few filled cells, all upper case (e.g. "TRIAL CHARACTERISTICS")."""
    filled = [str(c).strip() for c in cells if c not in (None, "")]
    return bool(filled) and len(filled) <= max(2, len(cells) // 4) and all(c.isupper() for c in filled)


def _rows(path: Path) -> List[List[Any]]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8") as handle:
            return [list(r) for r in csv.reader(handle)]
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return [list(r) for r in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_sheet(path: Path | str) -> Sheet:
    rows = [r for r in _rows(Path(path)) if any(c not in (None, "") for c in r)]
    if not rows:
        raise ValueError(f"{path}: no rows")
    header_index = 1 if len(rows) > 1 and _is_section_row(rows[0]) else 0
    raw_headers = rows[header_index]
    keep = [
        i for i, h in enumerate(raw_headers)
        if h not in (None, "") and not str(h).strip().startswith("_SYS") and not str(h).strip().lower().startswith("unnamed")
    ]
    headers = [str(raw_headers[i]).strip() for i in keep]
    if len(set(headers)) != len(headers):
        dupes = sorted({h for h in headers if headers.count(h) > 1})
        raise ValueError(f"{path}: duplicate headers {dupes}")
    data = []
    for raw in rows[header_index + 1:]:
        row = {headers[j]: clean_value(raw[i] if i < len(raw) else None) for j, i in enumerate(keep)}
        if any(v for k, v in row.items() if k != DOC_COLUMN):
            data.append(row)
    return Sheet(headers=headers, rows=data)


def write_sheet(path: Path | str, headers: List[str], rows: List[Dict[str, str]]) -> Path:
    """Write headers and rows as an .xlsx (or .csv) spreadsheet — used to build role-play inputs."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            for row in rows:
                writer.writerow([row.get(h, "") for h in headers])
        return path
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(h, "") for h in headers])
    workbook.save(path)
    return path
